from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from docx import Document
from docx.oxml.ns import qn

path = Path(r"C:\rafiki\outputs\rafiki-kkday-design-brief-seo-aeo-geo-audit.docx")
doc = Document(path)

text = "\n".join(p.text for p in doc.paragraphs)
text += "\n" + "\n".join(
    cell.text
    for table in doc.tables
    for row in table.rows
    for cell in row.cells
)
required = [
    "Approved scope addendum",
    "Live SEO/AEO/GEO Audit",
    "3,249",
    "2,740",
    "509 exact duplicates",
    "Structured Data Architecture",
    "Five Global Languages",
    "Simplified Chinese",
    "Hindi",
    "Spanish",
    "Arabic",
    "SEO/AEO/GEO Tools & Guides",
    "Additional Definition of Done",
]
missing = [value for value in required if value not in text]

hyperlinks = [
    rel
    for rel in doc.part.rels.values()
    if rel.reltype.endswith("/hyperlink")
]

geometry_errors = []
for table_index, table in enumerate(doc.tables[5:], start=5):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    grid_widths = [
        int(col.get(qn("w:w")))
        for col in table._tbl.tblGrid.findall(qn("w:gridCol"))
    ]
    declared = int(tbl_w.get(qn("w:w"))) if tbl_w is not None else -1
    if declared != sum(grid_widths):
        geometry_errors.append(
            f"table {table_index}: tblW={declared}, grid={sum(grid_widths)}"
        )
    for row_index, row in enumerate(table.rows):
        cell_widths = []
        for cell in row.cells:
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            cell_widths.append(int(tc_w.get(qn("w:w"))) if tc_w is not None else -1)
        if cell_widths != grid_widths:
            geometry_errors.append(
                f"table {table_index} row {row_index}: cells={cell_widths}, grid={grid_widths}"
            )

xml_errors = []
with ZipFile(path) as archive:
    for name in archive.namelist():
        if name.endswith((".xml", ".rels")):
            try:
                ET.fromstring(archive.read(name))
            except ET.ParseError as exc:
                xml_errors.append(f"{name}: {exc}")

new_bullets = 0
for paragraph in doc.paragraphs:
    if not paragraph.text:
        continue
    p_pr = paragraph._p.pPr
    if p_pr is None:
        continue
    num_pr = p_pr.find(qn("w:numPr"))
    p_style = p_pr.find(qn("w:pStyle"))
    if num_pr is not None and p_style is not None and p_style.get(qn("w:val")) == "ListParagraph":
        new_bullets += 1

print(f"file={path}")
print(f"bytes={path.stat().st_size}")
print(f"paragraphs={len(doc.paragraphs)}")
print(f"tables={len(doc.tables)}")
print(f"hyperlinks={len(hyperlinks)}")
print(f"numbered_list_paragraphs={new_bullets}")
print(f"missing_required={missing}")
print(f"geometry_errors={geometry_errors[:10]}")
print(f"xml_errors={xml_errors}")

if missing or geometry_errors or xml_errors:
    raise SystemExit(1)
