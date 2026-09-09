from pathlib import Path
from docx import Document

source = Path(r"C:\Users\Administrator\Downloads\rafiki-kkday-design-brief.docx")
doc = Document(source)

print(f"FILE: {source}")
print(f"PARAGRAPHS: {len(doc.paragraphs)}")
print(f"TABLES: {len(doc.tables)}")
print(f"SECTIONS: {len(doc.sections)}")
for si, section in enumerate(doc.sections):
    print(
        f"SECTION {si}: page={section.page_width.inches:.2f}x{section.page_height.inches:.2f}in "
        f"margins={section.top_margin.inches:.2f}/{section.right_margin.inches:.2f}/"
        f"{section.bottom_margin.inches:.2f}/{section.left_margin.inches:.2f}in"
    )

print("\nKEY STYLES")
for style_name in ["Normal", "Heading 1", "Heading 2", "Heading 3", "List Paragraph", "List Bullet", "List Number"]:
    try:
        style = doc.styles[style_name]
    except KeyError:
        print(f"{style_name}: MISSING")
        continue
    font = style.font
    pf = style.paragraph_format
    color = font.color.rgb if font.color and font.color.rgb else None
    print(
        f"{style_name}: font={font.name} size={font.size.pt if font.size else None} "
        f"bold={font.bold} color={color} before={pf.space_before.pt if pf.space_before else None} "
        f"after={pf.space_after.pt if pf.space_after else None} "
        f"line={pf.line_spacing}"
    )
print("\nPARAGRAPHS")
for i, paragraph in enumerate(doc.paragraphs):
    text = " ".join(paragraph.text.split())
    if text:
        style_name = paragraph.style.name if paragraph.style is not None else "(no style)"
        print(f"{i:04d}\t{style_name}\t{text}")

print("\nTABLES")
for ti, table in enumerate(doc.tables):
    print(f"TABLE {ti}: {len(table.rows)} rows x {len(table.columns)} cols, style={table.style.name if table.style else ''}")
    for ri, row in enumerate(table.rows):
        cells = [" ".join(cell.text.split()) for cell in row.cells]
        print(f"  {ri:03d}\t" + " | ".join(cells))
